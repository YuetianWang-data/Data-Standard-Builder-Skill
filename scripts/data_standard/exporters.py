# Writes validated standards to CSV and formatted Excel workbooks.

import csv

from .common import clean_text
from .tabular import require_openpyxl


BUSINESS_HEADERS = [
    ("business_standard_id", "Business Standard ID"),
    ("field_ids", "Field IDs"),
    ("table_names", "Table Names"),
    ("field_names", "Field Names"),
    ("field_chinese_names", "Field Descriptions"),
    ("standard_type", "Standard Type"),
    ("standard_code", "Standard or Rule Code"),
    ("standard_name", "Standard Name"),
    ("business_rule", "Business Rule"),
    ("data_type", "Data Type"),
    ("data_format", "Data Format"),
    ("value_domain", "Value Domain"),
    ("source", "Source"),
    ("source_location", "Source Location"),
    ("confidence", "Confidence"),
    ("reason", "Match or Recommendation Reason"),
    ("status", "Status"),
]

FIELD_HEADERS = [
    ("field_id", "Field ID"),
    ("table_name", "Table Name"),
    ("table_comment", "Table Description"),
    ("field_name", "Field Name"),
    ("field_comment", "Field Description"),
    ("data_type", "Data Type"),
    ("length", "Length"),
    ("nullable", "Nullable"),
    ("primary_key", "Primary Key"),
    ("default", "Default"),
    ("source_file", "Source File"),
    ("source_location", "Source Location"),
]

STANDARD_HEADERS = [
    ("standard_id", "Standard ID"),
    ("standard_code", "Data Element Code"),
    ("chinese_name", "Source-Language Data Element Name"),
    ("english_name", "English Data Element Name"),
    ("definition", "Definition"),
    ("data_type", "Data Type"),
    ("data_format", "Data Format"),
    ("value_domain", "Value Domain"),
    ("unit", "Unit"),
    ("source_file", "Source File"),
    ("source_location", "Source Location"),
    ("page", "Page"),
    ("extraction_method", "Extraction Method"),
]

RULE_HEADERS = [
    ("rule_id", "Rule ID"),
    ("system_name", "System Name"),
    ("table_name", "Table Name"),
    ("model_name", "Model Name"),
    ("rule_name", "Rule Name"),
    ("rule_content", "Rule Content"),
    ("problem_description", "Problem Description"),
    ("rule_status", "Rule Status"),
    ("scope", "Scope"),
    ("target", "Target"),
    ("severity", "Severity"),
    ("source", "Source"),
    ("source_location", "Source Location"),
]


def write_csv(path, rows, headers):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        labels = []
        for key, label in headers:
            labels.append(label)
        writer = csv.DictWriter(handle, fieldnames=labels, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            output_row = {}
            for key, label in headers:
                output_row[label] = row.get(key, "")
            writer.writerow(output_row)


def add_sheet(workbook, title, rows, headers):
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet = workbook.create_sheet(title)
    labels = []
    for key, label in headers:
        labels.append(label)
    sheet.append(labels)
    for row in rows:
        values = []
        for key, label in headers:
            values.append(row.get(key, ""))
        sheet.append(values)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        longest = 8
        for cell in column_cells[:200]:
            value_length = len(clean_text(cell.value))
            longest = max(longest, value_length)
        width = min(longest + 2, 50)
        sheet.column_dimensions[column_cells[0].column_letter].width = max(width, 10)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_workbook(path, business_rows, fields, standards, rules, uncovered, issues):
    openpyxl = require_openpyxl()
    workbook = openpyxl.Workbook()
    default = workbook.active
    workbook.remove(default)
    add_sheet(workbook, "Business Standards", business_rows, BUSINESS_HEADERS)
    add_sheet(workbook, "Database Fields", fields, FIELD_HEADERS)
    add_sheet(workbook, "Normative Standards", standards, STANDARD_HEADERS)
    add_sheet(workbook, "User Rules", rules, RULE_HEADERS)
    issue_rows = []
    for field_id in uncovered:
        issue_rows.append({"type": "Uncovered field", "detail": field_id})
    for issue in issues:
        issue_rows.append({"type": "Review note", "detail": issue})
    add_sheet(
        workbook,
        "Open Issues",
        issue_rows,
        [("type", "Issue Type"), ("detail", "Issue Details")],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
